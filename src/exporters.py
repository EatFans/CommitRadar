# -*- coding: utf-8 -*-
from typing import Dict, Any, List
import os
import csv
import json
import datetime
import logging
from pathlib import Path

# 尝试导入Excel相关库
try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

# 配置日志
logger = logging.getLogger('exporters')

class CommitDataExporter:
    """提交数据导出工具"""
    
    def __init__(self, export_dir: str = 'exports'):
        """
        初始化导出工具
        
        Args:
            export_dir: 导出文件保存目录
        """
        self.export_dir = export_dir
        
        # 确保导出目录存在
        os.makedirs(export_dir, exist_ok=True)
    
    def export_to_json(self, data: List[Dict[str, Any]], filename: str = None) -> str:
        """
        将提交数据导出为JSON文件
        
        Args:
            data: 提交数据列表
            filename: 文件名（可选）
            
        Returns:
            导出文件的路径
        """
        if not filename:
            timestamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f"github_commits_{timestamp}.json"
        
        filepath = os.path.join(self.export_dir, filename)
        
        with open(filepath, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        
        logger.info(f"数据已导出为JSON: {filepath}")
        return filepath
    
    def export_to_csv(self, data: List[Dict[str, Any]], filename: str = None) -> str:
        """
        将提交数据导出为CSV文件
        
        Args:
            data: 提交数据列表
            filename: 文件名（可选）
            
        Returns:
            导出文件的路径
        """
        if not filename:
            timestamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f"github_commits_{timestamp}.csv"
        
        filepath = os.path.join(self.export_dir, filename)
        
        with open(filepath, 'w', encoding='utf-8', newline='') as f:
            writer = csv.writer(f)
            
            # 写入表头
            writer.writerow(['用户名', '日期', '提交次数', '仓库', '提交标题', '提交时间', '提交URL'])
            
            # 写入数据
            for day_data in data:
                username = day_data['username']
                date = day_data['date']
                count = day_data['count']
                
                if day_data['commits']:
                    for commit in day_data['commits']:
                        writer.writerow([
                            username,
                            date,
                            count,
                            commit.get('repo', ''),
                            commit.get('title', ''),
                            commit.get('time', ''),
                            commit.get('url', '')
                        ])
                else:
                    # 如果没有提交详情，只写入基本信息
                    writer.writerow([username, date, count, '', '', '', ''])
        
        logger.info(f"数据已导出为CSV: {filepath}")
        return filepath
    
    def export_to_excel(self, data: List[Dict[str, Any]], filename: str = None) -> str:
        """
        将提交数据导出为Excel文件
        
        Args:
            data: 提交数据列表
            filename: 文件名（可选）
            
        Returns:
            导出文件的路径
        """
        if not EXCEL_AVAILABLE:
            logger.warning("无法导出为Excel: 缺少openpyxl库。请使用 'pip install openpyxl' 安装。")
            return self.export_to_csv(data, filename)
        
        if not filename:
            timestamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f"github_commits_{timestamp}.xlsx"
        
        filepath = os.path.join(self.export_dir, filename)
        
        # 创建工作簿和工作表
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "提交数据"
        
        # 设置表头
        headers = ['用户名', '日期', '提交次数', '仓库', '提交标题', '提交时间', '提交URL']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')
            cell.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
        
        # 写入数据
        row = 2
        for day_data in data:
            username = day_data['username']
            date = day_data['date']
            count = day_data['count']
            
            if day_data['commits']:
                for commit in day_data['commits']:
                    ws.cell(row=row, column=1, value=username)
                    ws.cell(row=row, column=2, value=date)
                    ws.cell(row=row, column=3, value=count)
                    ws.cell(row=row, column=4, value=commit.get('repo', ''))
                    ws.cell(row=row, column=5, value=commit.get('title', ''))
                    ws.cell(row=row, column=6, value=commit.get('time', ''))
                    
                    # 设置URL为超链接
                    url = commit.get('url', '')
                    if url:
                        ws.cell(row=row, column=7, value=url)
                        ws.cell(row=row, column=7).hyperlink = url
                    
                    row += 1
            else:
                # 如果没有提交详情，只写入基本信息
                ws.cell(row=row, column=1, value=username)
                ws.cell(row=row, column=2, value=date)
                ws.cell(row=row, column=3, value=count)
                row += 1
        
        # 调整列宽
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 20
        
        # 创建汇总工作表
        self._create_summary_sheet(wb, data)
        
        # 保存工作簿
        wb.save(filepath)
        
        logger.info(f"数据已导出为Excel: {filepath}")
        return filepath
    
    def _create_summary_sheet(self, workbook, data: List[Dict[str, Any]]):
        """创建汇总工作表"""
        ws = workbook.create_sheet(title="汇总")
        
        # 设置表头
        headers = ['日期', '提交次数']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')
            cell.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
        
        # 写入数据
        for row, day_data in enumerate(data, 2):
            ws.cell(row=row, column=1, value=day_data['date'])
            ws.cell(row=row, column=2, value=day_data['count'])
        
        # 计算总提交次数
        total_row = len(data) + 2
        ws.cell(row=total_row, column=1, value="总计")
        ws.cell(row=total_row, column=1).font = Font(bold=True)
        
        # 使用SUM函数计算总数
        ws.cell(row=total_row, column=2, value=f"=SUM(B2:B{total_row-1})")
        ws.cell(row=total_row, column=2).font = Font(bold=True)
        
        # 调整列宽
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15